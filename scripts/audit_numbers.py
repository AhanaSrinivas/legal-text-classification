"""Audit every number in the documents against results/.

Checks README.md, reports/writeup.pdf (extracted text), reports/slides.pptx
(slide text and speaker notes), reports/VIVA_QA.md and reports/CONTRIBUTIONS.md.
A number passes if it is
  * an integer found in a results/ file (including rounded/truncated forms), or
  * a value derived from results/ by src/report_facts.py, or
  * a documented setting in report_facts.CONFIG whose literal is still present
    in its source file, or
  * in the short STRUCTURAL list below (section numbers, citation years, ...).
Anything else is reported and the script exits with status 1.
Decimals must match an exact display value; arbitrary rounding is not allowed.
This checks numeric vocabulary, not whether a value labels the right model.
Generated-block and results-integrity tests provide those additional checks.
Run: python scripts/audit_numbers.py
"""

import re
import sys
from pathlib import Path

from pptx import Presentation
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from src.report_facts import CONFIG, facts  # noqa: E402

DOCUMENTS = ["README.md", "reports/writeup.pdf", "reports/slides.pptx", "reports/VIVA_QA.md",
             "reports/CONTRIBUTIONS.md"]

# Numbers that are not results: why each is allowed.
STRUCTURAL = {
    **{str(i): "small integer (section, phase, slide, list or count word)" for i in range(0, 21)},
    "95": "confidence level of the bootstrap intervals",
    "2020": "citation year (Iyer; SCDB; LEGAL-BERT)",
    "2022": "citation year (LexGLUE)",
    "2110.00976": "arXiv identifier of the LexGLUE paper",
    "256": "SHA-256 hash name",
    "438": "size in MB of legal_bert_state.pt on Hugging Face (438,049,972 bytes)",
    "8501": "Streamlit default port",
    "3.13": "Python version used (also in results/dataset_info.json)",
}

# Identifiers containing digits that are names, not quantities.
IGNORED_PATTERNS = [
    r"UE24CS352A", r"CS229", r"SHA-256", r"arXiv:2110\.00976", r"localhost:8501",
    r"`[^`\n]*`",                      # inline code: file names, commands, flags
    r"\[\d+\]",                        # reference markers in the PDF
    r"Release 01",                    # SCDB release identifier
]

NUMBER = re.compile(
    r"(?<![\w.])-?(?:\d+(?:\.\d+)?e-?\d+|\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+(?:\.\d+)?)(?![\w])"
)


def _norm(token):
    token = token.replace(",", "").replace("−", "-")
    if "e" in token:
        return f"{float(token):g}"
    return token


def _forms(value):
    """Display forms of a results value: rounded/truncated at 0-4 decimals, and as %."""
    out = set()
    for x in (value, value * 100):
        for decimals in range(0, 5):
            out.add(f"{x:.{decimals}f}")
            scale = 10 ** decimals
            out.add(f"{int(x * scale) / scale:.{decimals}f}")
    out.add(f"{value:g}")
    return out


def allowed_values(*, strict_decimals=True):
    """Use loose forms only for integers; the old mode is for coverage measurement."""
    allowed = {}
    for path in sorted((ROOT / "results").rglob("*")):
        if path.suffix not in {".json", ".csv", ".md", ".txt"}:
            continue
        text = path.read_text(encoding="utf-8", errors="ignore")
        for token in NUMBER.findall(text):
            value = float(_norm(token))
            for form in _forms(value):
                if not strict_decimals or re.fullmatch(r"-?\d+", form):
                    allowed.setdefault(form, f"results/{path.relative_to(ROOT / 'results')}")
    for key, value in facts().items():
        if isinstance(value, str):
            for token in NUMBER.findall(value):
                allowed.setdefault(_norm(token), f"report_facts: {key}")
    for key, (value, source, _) in CONFIG.items():
        for token in NUMBER.findall(value):
            allowed.setdefault(_norm(token), f"{source} ({key})")
    for token, reason in STRUCTURAL.items():
        allowed.setdefault(token, reason)
    return allowed


def check_config_literals():
    """Every documented setting must still appear literally in its source file."""
    problems = []
    for key, (_, source, literal) in CONFIG.items():
        if literal not in (ROOT / source).read_text(encoding="utf-8"):
            problems.append(f"CONFIG[{key!r}]: {literal!r} not found in {source}")
    return problems


def document_text(relative):
    path = ROOT / relative
    if path.suffix == ".pdf":
        return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    if path.suffix == ".pptx":
        parts = []
        for number, slide in enumerate(Presentation(str(path)).slides, 1):
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text_frame.text)
                if shape.has_table:
                    parts.extend(cell.text for row in shape.table.rows for cell in row.cells)
            parts.append(slide.notes_slide.notes_text_frame.text)
        return "\n".join(parts)
    return path.read_text(encoding="utf-8")


def audit_text(text, allowed=None):
    """Return the count and unmatched values, also usable for negative controls."""
    if allowed is None:
        allowed = allowed_values()
    for pattern in IGNORED_PATTERNS:
        text = re.sub(pattern, " ", text)
    found, unmatched = 0, []
    for match in NUMBER.finditer(text):
        found += 1
        token = _norm(match.group())
        if token not in allowed:
            context = text[max(0, match.start() - 40):match.end() + 40].replace("\n", " ")
            unmatched.append(f"{match.group()!r} in '...{context}...'")
    return found, unmatched


def audit():
    allowed = allowed_values()
    return {relative: audit_text(document_text(relative), allowed) for relative in DOCUMENTS}


def main():
    problems = check_config_literals()
    report = audit()
    for relative, (found, unmatched) in report.items():
        status = "PASS" if not unmatched else "FAIL"
        print(f"{status}  {relative}: {found} numbers checked, {len(unmatched)} unmatched")
        for item in unmatched:
            print(f"      {item}")
    for problem in problems:
        print(f"FAIL  {problem}")
    if not problems:
        print(f"PASS  {len(CONFIG)} documented settings found in their source files")
    if problems or any(unmatched for _, unmatched in report.values()):
        sys.exit(1)


if __name__ == "__main__":
    main()
