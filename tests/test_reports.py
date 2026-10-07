"""Read-only checks for committed report artifacts and their numeric provenance."""

import re
from pathlib import Path

import pytest
from pptx import Presentation
from pypdf import PdfReader

from scripts import audit_numbers, make_readme_results, make_viva_qa
from src.report_facts import render


ROOT = Path(__file__).resolve().parents[1]


def test_writeup_page_limit_and_required_sections():
    pdf = PdfReader(ROOT / "reports/writeup.pdf")
    assert 0 < len(pdf.pages) <= 2
    text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    for section in ("Problem statement", "Dataset details", "Approach",
                    "Brief implementation overview", "Conclusions"):
        assert section in text, section


def test_slides_count_notes_and_team_placeholders():
    deck = Presentation(ROOT / "reports/slides.pptx")
    assert len(deck.slides) == 15
    for index, slide in enumerate(deck.slides, 1):
        assert slide.has_notes_slide, f"Slide {index} has no notes"
        assert slide.notes_slide.notes_text_frame.text.strip(), f"Slide {index} has empty notes"
    title = "\n".join(shape.text for shape in deck.slides[0].shapes if shape.has_text_frame)
    for member in (1, 2):
        assert f"[Member {member} Name / SRN]" in title


def test_readme_generated_blocks_are_current():
    text = (ROOT / "README.md").read_text(encoding="utf-8")
    assert make_readme_results.update(text) == text


def test_viva_is_fresh_render():
    expected = make_viva_qa.HEADER + render(make_viva_qa.TEMPLATE.read_text(encoding="utf-8"))
    assert make_viva_qa.OUTPUT.read_text(encoding="utf-8") == expected


def test_document_number_audit():
    assert not audit_numbers.check_config_literals()
    report = audit_numbers.audit()
    assert set(report) == set(audit_numbers.DOCUMENTS)
    for document, (count, unmatched) in report.items():
        assert count > 0, document
        assert not unmatched, f"{document}: {unmatched}"


@pytest.mark.parametrize("text", ["Macro F1: 0.6421", "Accuracy: 83.7%"])
def test_number_audit_rejects_fabricated_decimals(text):
    count, unmatched = audit_numbers.audit_text(text)
    assert count == 1
    assert len(unmatched) == 1


def test_number_audit_preserves_identifier_exemptions():
    assert audit_numbers.audit_text("Version 2020 Release 01; `fake_0.6421.py`") == (1, [])


def repository_tree_paths(text):
    """Parse siblings, nested directories and grouped filenames in README section twelve."""
    section = re.search(r"^## 12\. Repository structure\s*\n(.*?)(?=^## |\Z)",
                        text, flags=re.MULTILINE | re.DOTALL)
    assert section, "Repository structure section missing"
    block = re.search(r"```[^\n]*\n(.*?)```", section.group(1), flags=re.DOTALL)
    assert block, "Repository tree missing"
    parents = {-1: Path()}
    paths = []
    for line in block.group(1).splitlines()[1:]:
        match = re.fullmatch(r"((?:│   |    )*)(?:├── |└── )(.*)", line)
        assert match, f"Unrecognised tree entry: {line!r}"
        depth = len(match.group(1)) // 4
        # Descriptions follow the filename expression; comma and spaced-slash
        # separators join filenames, while ordinary slashes belong to paths.
        names = re.match(r"[^\s,]+(?:,\s+[^\s,]+| / [^\s,]+)*", match.group(2)).group()
        parent = parents[depth - 1]
        entries = re.split(r",\s*|\s+/\s+", names)
        paths.extend(parent / entry for entry in entries)
        if len(entries) == 1 and entries[0].endswith("/"):
            parents[depth] = parent / entries[0]
    assert paths, "Repository tree has no entries"
    return paths


def test_every_readme_tree_path_exists():
    paths = repository_tree_paths((ROOT / "README.md").read_text(encoding="utf-8"))
    missing = [str(path) for path in paths if not (ROOT / path).exists()]
    assert not missing, f"Missing README paths: {missing}"


def test_tree_parser_keeps_directory_context_and_grouped_files():
    text = """## 12. Repository structure
```
project/
├── README.md, PROJECT_PLAN.md   root files
├── reports/
│   ├── writeup.py / writeup.pdf   sources
│   └── templates/
│       └── VIVA_QA.template.md
└── app.py   application
```
## 13. Limitations
"""
    assert repository_tree_paths(text) == list(map(Path, [
        "README.md", "PROJECT_PLAN.md", "reports", "reports/writeup.py",
        "reports/writeup.pdf", "reports/templates", "reports/templates/VIVA_QA.template.md", "app.py",
    ]))
